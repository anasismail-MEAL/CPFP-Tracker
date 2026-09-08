"""Monthly aggregates and the XLSX export.

The per-focal-point figures deliberately reproduce the CPFP_Daily_Activity_Tracker
v4.0 `Monthly Summary` sheet, including its period rule (>= first of month, < first
of next month), so a CPFP can check the app against the workbook they already know.
The case figures are the part the workbook cannot produce.
"""

import statistics
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import db

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SECTION_FONT = Font(bold=True, color="1F4E79")

# Follow-ups the workbook counts as still owing action.
OPEN_FOLLOWUP = "followup = 'Yes' AND status NOT IN ('Closed', 'Not applicable')"


def period_bounds(year: int, month: int):
    start = date(year, month, 1)
    end = date(year + (month == 12), (month % 12) + 1, 1)
    return start.isoformat(), end.isoformat()


def _hours_between(a: str, b: str) -> float:
    return (datetime.fromisoformat(b) - datetime.fromisoformat(a)).total_seconds() / 3600


def per_focal_point(conn, start, end):
    rows = conn.execute(
        f"""
        SELECT u.id, u.full_name, u.agency, u.camp, u.role,
               COUNT(t.id)                    AS activities,
               COALESCE(SUM(t.n_children), 0) AS children,
               COALESCE(SUM(t.n_adults), 0)   AS adults,
               COALESCE(SUM(CASE WHEN {OPEN_FOLLOWUP} THEN 1 ELSE 0 END), 0) AS open_followups
        FROM users u
        LEFT JOIN tasks t ON t.user_id = u.id AND t.tdate >= ? AND t.tdate < ?
        WHERE u.active = 1
        GROUP BY u.id
        ORDER BY u.role DESC, u.full_name
        """,
        (start, end),
    ).fetchall()
    return [dict(r) for r in rows]


def breakdown(conn, start, end, column, kind):
    """Entries/children/adults per activity type, or entries per stakeholder."""
    ordered = db.ref_list(conn, kind)
    rows = conn.execute(
        f"""
        SELECT {column} AS k, COUNT(*) AS entries,
               COALESCE(SUM(n_children), 0) AS children,
               COALESCE(SUM(n_adults), 0)   AS adults
        FROM tasks WHERE tdate >= ? AND tdate < ? GROUP BY {column}
        """,
        (start, end),
    ).fetchall()
    found = {r["k"]: dict(r) for r in rows}
    out = []
    for key in ordered:
        hit = found.pop(key, None)
        out.append(hit or {"k": key, "entries": 0, "children": 0, "adults": 0})
    out.extend(found.values())  # anything recorded outside the agreed list stays visible
    return out


def case_stats(conn, start, end):
    opened = conn.execute(
        "SELECT COUNT(*) c FROM cases WHERE opened_at >= ? AND opened_at < ?", (start, end)
    ).fetchone()["c"]

    reunified = conn.execute(
        """SELECT opened_at, closed_at, camp_report FROM cases
           WHERE status = 'reunified' AND closed_at >= ? AND closed_at < ?""",
        (start, end),
    ).fetchall()

    still_open = conn.execute(
        "SELECT COUNT(*) c FROM cases WHERE status IN ('open', 'matched') AND opened_at < ?",
        (end,),
    ).fetchone()["c"]

    hours = [_hours_between(r["opened_at"], r["closed_at"]) for r in reunified]

    by_camp = {}
    for r in reunified:
        camp = r["camp_report"] or "(unknown)"
        by_camp[camp] = by_camp.get(camp, 0) + 1

    by_type = {
        r["ctype"]: r["c"]
        for r in conn.execute(
            """SELECT ctype, COUNT(*) c FROM cases
               WHERE opened_at >= ? AND opened_at < ? GROUP BY ctype""",
            (start, end),
        ).fetchall()
    }

    return {
        "opened": opened,
        "reunified": len(reunified),
        "still_open": still_open,
        "median_hours": round(statistics.median(hours), 1) if hours else None,
        # The ToR sets a 24-hour reunification target, so report against it.
        "within_24h": sum(1 for h in hours if h <= 24),
        "by_camp": sorted(by_camp.items(), key=lambda kv: (-kv[1], kv[0])),
        "missing": by_type.get("missing", 0),
        "found": by_type.get("found", 0),
    }


def monthly(conn, year, month):
    start, end = period_bounds(year, month)
    return {
        "year": year,
        "month": month,
        "start": start,
        "end": end,
        "people": per_focal_point(conn, start, end),
        "activities": breakdown(conn, start, end, "activity_type", "activity_type"),
        "stakeholders": breakdown(conn, start, end, "stakeholder", "stakeholder"),
        "cases": case_stats(conn, start, end),
        "totals": conn.execute(
            f"""SELECT COUNT(*) activities,
                       COALESCE(SUM(n_children), 0) children,
                       COALESCE(SUM(n_adults), 0)   adults,
                       COALESCE(SUM(CASE WHEN {OPEN_FOLLOWUP} THEN 1 ELSE 0 END), 0) open_followups
                FROM tasks WHERE tdate >= ? AND tdate < ?""",
            (start, end),
        ).fetchone(),
    }


# --------------------------------------------------------------------------- export

def _head(ws, row, labels, col=1):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center")


def to_xlsx(data) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Summary"
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 16

    ws["A1"] = "CPFP MONTHLY SUMMARY"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Child Protection Sub-Sector (CPSS) | Cox's Bazar | %d-%02d" % (
        data["year"],
        data["month"],
    )

    r = 4
    ws.cell(row=r, column=1, value="BY FOCAL POINT").font = SECTION_FONT
    r += 1
    _head(ws, r, ["Focal point", "Agency", "Camp", "Activities", "Children",
                  "Adults", "Open follow-ups"])
    for p in data["people"]:
        r += 1
        vals = [p["full_name"], p["agency"], p["camp"], p["activities"],
                p["children"], p["adults"], p["open_followups"]]
        for i, v in enumerate(vals):
            ws.cell(row=r, column=1 + i, value=v)

    t = data["totals"]
    r += 2
    ws.cell(row=r, column=1, value="OVERALL").font = SECTION_FONT
    r += 1
    _head(ws, r, ["Total activities", "Children involved", "Adults involved", "Open follow-ups"])
    r += 1
    for i, v in enumerate([t["activities"], t["children"], t["adults"], t["open_followups"]]):
        ws.cell(row=r, column=1 + i, value=v)

    r += 2
    ws.cell(row=r, column=1, value="ACTIVITY BREAKDOWN").font = SECTION_FONT
    r += 1
    _head(ws, r, ["Activity type", "Entries", "Children", "Adults"])
    for a in data["activities"]:
        r += 1
        for i, v in enumerate([a["k"], a["entries"], a["children"], a["adults"]]):
            ws.cell(row=r, column=1 + i, value=v)

    r += 2
    ws.cell(row=r, column=1, value="STAKEHOLDER / SECTOR BREAKDOWN").font = SECTION_FONT
    r += 1
    _head(ws, r, ["Stakeholder / sector", "Entries"])
    for s in data["stakeholders"]:
        r += 1
        ws.cell(row=r, column=1, value=s["k"])
        ws.cell(row=r, column=2, value=s["entries"])

    cs = data["cases"]
    ws2 = wb.create_sheet("Missing & Reunification")
    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 16
    ws2["A1"] = "MISSING / FOUND CHILD COORDINATION"
    ws2["A1"].font = TITLE_FONT
    ws2["A2"] = "Counts only. No identifying child data leaves this system."

    rows = [
        ("Cases opened", cs["opened"]),
        ("  of which missing-child reports", cs["missing"]),
        ("  of which found-child reports", cs["found"]),
        ("Children reunified", cs["reunified"]),
        ("Reunified within 24 hours (ToR target)", cs["within_24h"]),
        ("Median hours to reunification", cs["median_hours"]),
        ("Cases still open at period end", cs["still_open"]),
    ]
    r = 4
    _head(ws2, r, ["Indicator", "Value"])
    for label, val in rows:
        r += 1
        ws2.cell(row=r, column=1, value=label)
        ws2.cell(row=r, column=2, value="n/a" if val is None else val)

    r += 2
    ws2.cell(row=r, column=1, value="REUNIFICATIONS BY CAMP").font = SECTION_FONT
    r += 1
    _head(ws2, r, ["Camp", "Reunified"])
    for camp, n in cs["by_camp"]:
        r += 1
        ws2.cell(row=r, column=1, value=camp)
        ws2.cell(row=r, column=2, value=n)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
