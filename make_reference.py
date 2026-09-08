"""Regenerate reference_lists.json from the CPFP tracker workbook.

The workbook is the source of truth for the camp, activity, stakeholder and status
lists. It is not committed (it lives one directory up, alongside material that must
not be published), so this writes a JSON snapshot that a fresh clone or a deployment
can seed from. Re-run it whenever CPSS agrees a change to any of those lists.
"""

import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE.parent / "CPFP_Daily_Activity_Tracker_v4_0.xlsx"
OUT = HERE / "reference_lists.json"

# Dropdowns sheet: one column per list, header in row 1, values from row 2 down.
COLUMNS = {"camp": 1, "activity_type": 2, "stakeholder": 3, "status": 5}


def read_lists(path=WORKBOOK):
    ws = load_workbook(path, data_only=True)["Dropdowns"]
    lists = {}
    for kind, col in COLUMNS.items():
        vals = []
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is not None and str(v).strip():
                vals.append(str(v).strip())
        lists[kind] = vals
    return lists


if __name__ == "__main__":
    if not WORKBOOK.exists():
        raise SystemExit("Workbook not found: %s" % WORKBOOK)
    lists = read_lists()
    OUT.write_text(json.dumps(lists, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("Wrote %s from %s" % (OUT.name, WORKBOOK.name))
    for kind, vals in lists.items():
        print("  %-14s %d values" % (kind, len(vals)))
